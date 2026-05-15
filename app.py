import streamlit as st
import pandas as pd
import re
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from io import BytesIO

# Configuração visual da página
st.set_page_config(page_title="Automação Parvi", page_icon="📊")

st.title("📊 Extrator de Pendências Globus")
st.markdown("Arraste seu arquivo TXT aqui para gerar a planilha formatada.")

# Upload do arquivo via navegador
arquivo_txt = st.file_uploader("Selecione o arquivo TXT", type=['txt'])

if arquivo_txt is not None:
    # Lendo os dados
    linhas = arquivo_txt.getvalue().decode("latin-1").splitlines()
    total_linhas = len(linhas)
    
    dados_lista = []
    mat, nome, func = None, None, None
    
    # Barra de progresso visual (a equipe vai amar isso!)
    barra = st.progress(0)
    status = st.empty()

    for i, linha in enumerate(linhas, 1):
        if i % 100 == 0 or i == total_linhas:
            barra.progress(i / total_linhas)
            status.text(f"Analisando linha {i} de {total_linhas}...")

        # Captura cabeçalho do funcionário
        if re.match(r'^\d{6}/', linha):
            mat = linha[0:6].strip()
            nome = linha[14:35].strip()
            func = linha[35:50].strip()

        # Busca datas e ocorrências (Pega todos os dias do funcionário)
        m_data = re.search(r'(\d{2}/\d{2}/\d{4})\s+(.*)', linha)
        if m_data and mat:
            data_texto = m_data.group(1)
            conteudo = m_data.group(2).upper()
            
            if "**" in conteudo or "SEM MOVIMENTO" in conteudo or "ATESTA" in conteudo:
                dados_lista.append({
                    "Matricula": mat, "NOME": nome, "FUNÇÃO": func,
                    "Data": data_texto, "Globus": conteudo.replace("**", "").strip(),
                    "OBSERVAÇÕES": ""
                })

    if dados_lista:
        st.success(f"✅ Concluído! {len(dados_lista)} pendências encontradas.")
        
        # Organização dos dados
        df = pd.DataFrame(dados_lista)
        df['D_Ref'] = pd.to_datetime(df['Data'], format='%d/%m/%Y')
        df = df.sort_values(by=['NOME', 'D_Ref']).drop(columns=['D_Ref'])

        # Criando o Excel na memória do servidor
        output = BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "Sem Movimentos"
        ws.append(["Matricula", "NOME", "FUNÇÃO", "Data", "Globus", "OBSERVAÇÕES"])

        # Estilização Parvi (Cabeçalho escuro e bordas)
        preto_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
        branco_font = Font(color="FFFFFF", bold=True)
        borda_fina = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        for cell in ws[1]:
            cell.fill = preto_fill
            cell.font = branco_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = borda_fina

        for r_idx, row in enumerate(df.values, start=2):
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.border = borda_fina

        # Coluna de Obs em azul claro (estético)
        azul_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        for r in range(2, len(df) + 2):
            ws.cell(row=r, column=6).fill = azul_fill

        # Ajuste de larguras
        for col, larg in zip("ABCDEF", [12, 35, 25, 15, 25, 40]):
            ws.column_dimensions[col].width = larg

        wb.save(output)
        
        # Botão de download
        st.download_button(
            label="📥 Baixar Planilha Excel",
            data=output.getvalue(),
            file_name="Relatorio_Final_Parvi.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    else:
        st.warning("Nenhuma pendência encontrada com esses filtros.")
