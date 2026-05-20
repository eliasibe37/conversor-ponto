import streamlit as st
import pandas as pd
import io
from openpyxl.styles import Font, PatternFill, Alignment

# --- Lógica de Processamento ---
def processar_globus(arquivo_txt):
    linhas = arquivo_txt.getvalue().decode("latin-1").splitlines()
    dados = []
    mat = nome = func = None
    
    for linha in linhas:
        # Fatiamento fixo
        if len(linha) > 50 and linha[0:6].isdigit() and "/" in linha[0:13]:
            mat = linha[0:13].strip().split('/')[0]
            nome = linha[14:34].strip()
            func = linha[34:52].strip()
        
        if "** SEM MOVIMENTO **" in linha:
            data = linha[50:62].strip()
            if mat:
                dados.append({
                    "Matricula": mat,
                    "NOME": nome,
                    "FUNÇÃO": func,
                    "Data": data,
                    "Globus": "SEM MOVIMENTO",
                    "OBSERVAÇÕES": ""
                })
    return pd.DataFrame(dados)

# --- Lógica de Estilização ---
def gerar_excel_estilizado(df):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name="Pendencias")
        ws = writer.sheets["Pendencias"]
        titulo_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        titulo_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = titulo_fill
            cell.font = titulo_font
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 18
    return output.getvalue()

# --- Interface ---
st.title("📊 Listage De Sem Movimentos")
arquivo = st.file_uploader("Suba o arquivo .txt", type=['txt'])

# Esta checagem garante que o código só rode se o arquivo existir
if arquivo is not None:
    # 1. Processa
    df = processar_globus(arquivo)
    
    # 2. Exibe métrica (Quantidade)
    st.metric(label="Total de Ocorrências Encontradas", value=len(df))
    
    # 3. Exibe a tabela (Prévia)
    st.subheader("Pré-visualização dos Dados")
    st.dataframe(df, use_container_width=True)
    
    # 4. Gera botão de download
    excel_data = gerar_excel_estilizado(df)
    st.download_button(
        label="📥 Baixar Excel Estilizado", 
        data=excel_data, 
        file_name="Relatorio_Processado.xlsx"
    )
else:
    st.info("Por favor, selecione um arquivo .txt para começar.")
