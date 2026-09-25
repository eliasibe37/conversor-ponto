import streamlit as st
import pandas as pd
import io
import re
import plotly.express as px
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Comparativo Executivo Globus", page_icon="📊", layout="wide")

st.title("📊 Comparativo De Resumo Do Acumulado - TXT Globus")
st.markdown("`Compare Até 5 Meses Simultaneamente` - Analise De Resumo.")

# ==========================================
# FUNÇÕES DE TRATAMENTO DE DADOS
# ==========================================
def converter_horas_para_decimal(valor_num):
    """Converte formato relógio (HH.MM) para horas decimais reais (HH.DEC)."""
    horas = int(valor_num)
    minutos = round((valor_num - horas) * 100)
    return horas + (minutos / 60.0)

def extrair_dados_globus_completo(txt_content, rotulo_mes):
    linhas = txt_content.splitlines()
    registros_finais = []
    totais_rodape = {}
    
    matricula_atual = None
    nome_atual = None
    funcao_atual = None
    dentro_do_resumo_final = False

    EVENTOS_NAO_HORAS = ["00001", "1", "00002", "2", "SALARIO BASE", "DIAS", "ATESTADO"]

    for linha in linhas:
        if not linha.strip() or "-------" in linha:
            continue
            
        linha_upper = linha.upper()
        
        if "RESUMO POR EVENTO" in linha_upper:
            dentro_do_resumo_final = True
            continue
            
        # 1. PROCESSAMENTO DO RODAPÉ OFICIAL
        if dentro_do_resumo_final:
            partes_rodape = [p.strip() for p in linha.split("  ") if p.strip()]
            if len(partes_rodape) >= 3:
                cod = partes_rodape[0]
                if len(cod) == 5 and cod.isdigit():
                    val_str = partes_rodape[-1]
                    desc = " ".join(partes_rodape[1:-1]).strip()
                    if not any(x in desc.upper() for x in ["TOTAL", "RESUMO", "EVENTO", "REFERENCIA"]):
                        try:
                            val_original = float(val_str.replace('.', '').replace(',', '.'))
                            if cod not in EVENTOS_NAO_HORAS and not any(e in desc.upper() for e in EVENTOS_NAO_HORAS):
                                val_final = round(converter_horas_para_decimal(val_original), 2)
                            else:
                                val_final = round(val_original, 2)

                            totais_rodape[cod] = {
                                "MÊS": rotulo_mes,
                                "CÓDIGO": cod,
                                "EVENTO": desc,
                                "TOTAL REFERÊNCIA": val_final
                            }
                        except:
                            pass
            continue

        if any(termo in linha_upper for termo in ["PERIODO:", "REGISTRO NOME", "PAGINA :", "EMPRESA", "CNPJ:"]):
            continue

        # 2. CAPTURA DO FUNCIONÁRIO
        inicio_linha = linha[:6].strip()
        if inicio_linha.isdigit() and len(inicio_linha) >= 4 and "TOTAL" not in linha_upper:
            matricula_atual = inicio_linha
            nome_atual = linha[9:40].strip()
            funcao_atual = linha[40:57].strip()
            if "/" in nome_atual:
                nome_atual = nome_atual.split()[0]

        # 3. EXTRAÇÃO DOS EVENTOS DETALHADOS
        if len(linha) > 70:
            metade_direita = linha[70:]
            match_evento = re.search(r'(\d{5})\s+([A-Z0-9\.\s%/-]+?)\s+([\d\.,\-]+)$', metade_direita.upper())
            
            if match_evento:
                trecho_cod_evento = match_evento.group(1).strip()
                nome_evento = match_evento.group(2).strip()
                referencia_str = match_evento.group(3).strip()
                
                if "TOTAL" in nome_evento:
                    continue
                    
                if matricula_atual and trecho_cod_evento:
                    try:
                        ref_bruta = float(referencia_str.replace('.', '').replace(',', '.'))
                        if trecho_cod_evento not in EVENTOS_NAO_HORAS and "SALARIO" not in nome_evento:
                            ref_final = round(converter_horas_para_decimal(ref_bruta), 2)
                        else:
                            ref_final = round(ref_bruta, 2)
                        
                        registros_finais.append({
                            "MÊS": rotulo_mes,
                            "MATRICULA": matricula_atual,
                            "NOME": nome_atual,
                            "FUNÇÃO": funcao_atual,
                            "COD DO EVENTO": trecho_cod_evento,
                            "EVENTO": nome_evento,
                            "REFERENCIA": ref_final
                        })
                    except:
                        pass

    df_detalhado = pd.DataFrame(registros_finais)
    
    if totais_rodape:
        df_resumo = pd.DataFrame(totais_rodape.values())
    else:
        df_resumo = df_detalhado.groupby(["COD DO EVENTO", "EVENTO"], as_index=False)["REFERENCIA"].sum()
        df_resumo.columns = ["CÓDIGO", "EVENTO", "TOTAL REFERÊNCIA"]
        df_resumo["MÊS"] = rotulo_mes
        df_resumo["TOTAL REFERÊNCIA"] = df_resumo["TOTAL REFERÊNCIA"].round(2)
        
    return df_detalhado, df_resumo


# ==========================================
# PAINEL LATERAL: UPLOAD MULTI-MÊS (2 A 5)
# ==========================================
st.sidebar.header("📁 Carregamento dos Arquivos")
qtd_meses = st.sidebar.number_input("Quantos meses deseja comparar?", min_value=2, max_value=5, value=2, step=1)

arquivos_carregados = []
meses_nomes = []

for i in range(int(qtd_meses)):
    st.sidebar.subheader(f"Mês {i+1}")
    nome_mes = st.sidebar.text_input(f"Rótulo do Mês {i+1} (ex: Jan/26)", value=f"Mês {i+1}", key=f"mes_nome_{i}")
    file_txt = st.sidebar.file_uploader(f"TXT do Globus ({nome_mes})", type=["txt"], key=f"file_{i}")
    
    if file_txt is not None:
        arquivos_carregados.append((nome_mes, file_txt))
        meses_nomes.append(nome_mes)

# ==========================================
# PROCESSAMENTO CONJUNTO
# ==========================================
if len(arquivos_carregados) >= 2:
    todos_detalhes = []
    todos_resumos = []

    for nome_m, arq in arquivos_carregados:
        bytes_data = arq.getvalue()
        try:
            txt_content = bytes_data.decode("utf-8")
        except UnicodeDecodeError:
            txt_content = bytes_data.decode("latin1", errors="ignore")
            
        df_det, df_res = extrair_dados_globus_completo(txt_content, rotulo_mes=nome_m)
        todos_detalhes.append(df_det)
        todos_resumos.append(df_res)

    df_detalhado_geral = pd.concat(todos_detalhes, ignore_index=True)
    df_resumo_geral = pd.concat(todos_resumos, ignore_index=True)

    # MATRIZ COMPARATIVA (Pivoteada)
    df_pivot = df_resumo_geral.pivot_table(
        index=["CÓDIGO", "EVENTO"], 
        columns="MÊS", 
        values="TOTAL REFERÊNCIA", 
        aggfunc="sum"
    ).fillna(0).reset_index()

    colunas_ordenadas = ["CÓDIGO", "EVENTO"] + meses_nomes
    df_pivot = df_pivot[colunas_ordenadas]

    # COMPARAÇÃO PRINCIPAL: ÚLTIMO MÊS VS MÊS ANTERIOR
    penultimo_mes = meses_nomes[-2]
    ultimo_mes = meses_nomes[-1]

    df_pivot[f"Var. Absoluta ({ultimo_mes} vs {penultimo_mes})"] = (df_pivot[ultimo_mes] - df_pivot[penultimo_mes]).round(2)
    var_perc = ((df_pivot[ultimo_mes] - df_pivot[penultimo_mes]) / df_pivot[penultimo_mes].replace(0, pd.NA)) * 100
    df_pivot["Var. % Recente"] = var_perc.fillna(0).astype(float).round(2)

    col_var_abs = f"Var. Absoluta ({ultimo_mes} vs {penultimo_mes})"

    def classificar_status(row):
        diff = row[col_var_abs]
        if diff > 0.05:
            return "▲ Aumentou"
        elif diff < -0.05:
            return "▼ Diminuiu"
        else:
            return "▬ Estável"

    df_pivot["Tendência Recente"] = df_pivot.apply(classificar_status, axis=1)

    # ==========================================
    # PAINEL DE INSIGHTS NO STREAMLIT
    # ==========================================
    st.subheader("💡 Resumo de Variações Recentes (Último Mês vs. Penúltimo Mês)")
    st.markdown(f"Análise focada na mudança de **{penultimo_mes}** para **{ultimo_mes}**:")

    maiores_aumentos = df_pivot[df_pivot[col_var_abs] > 0].sort_values(by=col_var_abs, ascending=False).head(5)
    maiores_quedas = df_pivot[df_pivot[col_var_abs] < 0].sort_values(by=col_var_abs, ascending=True).head(5)

    c1, c2 = st.columns(2)
    with c1:
        st.error(f"🚨 Maior Crescimento Recente ({ultimo_mes} vs {penultimo_mes})")
        if not maiores_aumentos.empty:
            for _, row in maiores_aumentos.iterrows():
                st.write(f"• **{row['EVENTO']}**: +{row[col_var_abs]:,.2f} ({row['Var. % Recente']:.2f}%) — Subiu de {row[penultimo_mes]:,.2f} para {row[ultimo_mes]:,.2f}")
        else:
            st.write("Nenhum evento teve crescimento relevante.")

    with c2:
        st.success(f"✅ Maior Redução Recente ({ultimo_mes} vs {penultimo_mes})")
        if not maiores_quedas.empty:
            for _, row in maiores_quedas.iterrows():
                st.write(f"• **{row['EVENTO']}**: {row[col_var_abs]:,.2f} ({row['Var. % Recente']:.2f}%) — Caiu de {row[penultimo_mes]:,.2f} para {row[ultimo_mes]:,.2f}")
        else:
            st.write("Nenhum evento teve redução relevante.")

    st.markdown("---")

    # ==========================================
    # DASHBOARD STREAMLIT
    # ==========================================
    st.subheader("📊 Visualização de Tendências Temporais")
    col_dash1, col_dash2 = st.columns([1, 2])

    with col_dash1:
        evento_selecionado = st.selectbox("Selecione um Evento para analisar a curva completa:", df_pivot["EVENTO"].unique())
        dados_grafico = df_resumo_geral[df_resumo_geral["EVENTO"] == evento_selecionado]
        fig = px.line(
            dados_grafico, 
            x="MÊS", 
            y="TOTAL REFERÊNCIA", 
            text="TOTAL REFERÊNCIA",
            markers=True,
            title=f"Evolução Histórica: {evento_selecionado}"
        )
        fig.update_traces(textposition="top center")
        st.plotly_chart(fig)

    with col_dash2:
        st.write(f"### Top Eventos em {ultimo_mes}")
        top_eventos = df_resumo_geral[df_resumo_geral["MÊS"] == ultimo_mes].sort_values(by="TOTAL REFERÊNCIA", ascending=False).head(8)
        fig_bar = px.bar(
            top_eventos, 
            x="TOTAL REFERÊNCIA", 
            y="EVENTO", 
            orientation='h',
            text="TOTAL REFERÊNCIA",
            color="TOTAL REFERÊNCIA",
            title=f"Eventos acumulados ({ultimo_mes})"
        )
        fig_bar.update_layout(yaxis={'categoryorder':'total ascending'})
        st.plotly_chart(fig_bar)

    st.subheader("📋 Tabela Comparativa Detalhada (Histórico Completo)")
    st.dataframe(
        df_pivot, 
        column_config={
            "Var. % Recente": st.column_config.NumberColumn("Var. % Recente", format="%.2f %%"),
        }
    )

    # ==========================================
    # GERAÇÃO DO EXCEL EXECUTIVO
    # ==========================================
    buffer = io.BytesIO()
    wb = openpyxl.Workbook()

    NAVY_DARK = "1B365D"
    BLUE_HEADER = "2C4D75"
    GRAY_CARD = "EAEFF5"
    RED_SOFT = "FADBD8"
    RED_TEXT = "78281F"
    GREEN_SOFT = "D4EFDF"
    GREEN_TEXT = "145A32"
    GRAY_SOFT = "E5E7E9"
    GRAY_TEXT = "515A5A"
    BORDER_COLOR = "D0D7DE"

    font_title = Font(name="Segoe UI", size=16, bold=True, color=NAVY_DARK)
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="555555")
    font_section = Font(name="Segoe UI", size=12, bold=True, color=NAVY_DARK)
    font_head = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_cell = Font(name="Segoe UI", size=10)

    font_red = Font(name="Segoe UI", size=10, bold=True, color=RED_TEXT)
    font_green = Font(name="Segoe UI", size=10, bold=True, color=GREEN_TEXT)
    font_gray = Font(name="Segoe UI", size=10, bold=True, color=GRAY_TEXT)

    fill_red = PatternFill(start_color=RED_SOFT, end_color=RED_SOFT, fill_type="solid")
    fill_green = PatternFill(start_color=GREEN_SOFT, end_color=GREEN_SOFT, fill_type="solid")
    fill_gray = PatternFill(start_color=GRAY_SOFT, end_color=GRAY_SOFT, fill_type="solid")
    fill_header = PatternFill(start_color=BLUE_HEADER, end_color=BLUE_HEADER, fill_type="solid")
    fill_card = PatternFill(start_color=GRAY_CARD, end_color=GRAY_CARD, fill_type="solid")
    
    border_light = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    # ABA 1: EXECUTIVE SUMMARY
    ws_dash = wb.active
    ws_dash.title = "Executive Summary"
    ws_dash.views.sheetView[0].showGridLines = True

    ws_dash.cell(row=2, column=2, value="RELATÓRIO EXECUTIVO DE FREQUÊNCIA E OCORRÊNCIAS").font = font_title
    ws_dash.cell(row=3, column=2, value=f"Análise Focada na Variação Recente: {penultimo_mes} ➔ {ultimo_mes} | Base Globus").font = font_subtitle

    # Cards KPI
    tot_penultimo = df_pivot[penultimo_mes].sum()
    tot_ultimo = df_pivot[ultimo_mes].sum()
    var_rec_abs = tot_ultimo - tot_penultimo
    var_rec_pct = (var_rec_abs / tot_penultimo * 100) if tot_penultimo > 0 else 0

    ws_dash.merge_cells("B5:D6")
    c1 = ws_dash.cell(row=5, column=2)
    c1.value = f"TOTAL EM {ultimo_mes.upper()}\n{tot_ultimo:,.2f} Horas/Ocorr."
    c1.font = Font(name="Segoe UI", size=11, bold=True, color=NAVY_DARK)
    c1.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for r in range(5, 7):
        for c in range(2, 5):
            ws_dash.cell(row=r, column=c).fill = fill_card
            ws_dash.cell(row=r, column=c).border = border_light

    ws_dash.merge_cells("E5:G6")
    c2 = ws_dash.cell(row=5, column=5)
    c2.value = f"VARIAÇÃO RECENTE ({ultimo_mes} vs {penultimo_mes})\n{var_rec_abs:+,.2f} ({var_rec_pct:+.2f}%)"
    c2.font = Font(name="Segoe UI", size=11, bold=True, color="900C3F" if var_rec_abs > 0 else "1E8449")
    c2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for r in range(5, 7):
        for c in range(5, 8):
            ws_dash.cell(row=r, column=c).fill = fill_card
            ws_dash.cell(row=r, column=c).border = border_light

    # Tabela: Maiores Aumentos Recentes
    ws_dash.cell(row=8, column=2, value=f"🚨 TOP AUMENTOS NO ÚLTIMO MÊS ({penultimo_mes} ➔ {ultimo_mes})").font = font_section
    headers_resumo = ["CÓDIGO", "DESCRICÃO DO EVENTO", penultimo_mes, ultimo_mes, "VARIAÇÃO ABSOLUTA", "VARIAÇÃO %"]
    
    for idx, h in enumerate(headers_resumo, start=2):
        cell = ws_dash.cell(row=9, column=idx, value=h)
        cell.fill = fill_header
        cell.font = font_head
        cell.alignment = Alignment(horizontal="center", vertical="center")

    curr_r = 10
    for _, r in maiores_aumentos.iterrows():
        ws_dash.cell(row=curr_r, column=2, value=r["CÓDIGO"]).alignment = Alignment(horizontal="center")
        ws_dash.cell(row=curr_r, column=3, value=r["EVENTO"])
        ws_dash.cell(row=curr_r, column=4, value=r[penultimo_mes]).number_format = "#,##0.00"
        ws_dash.cell(row=curr_r, column=5, value=r[ultimo_mes]).number_format = "#,##0.00"
        ws_dash.cell(row=curr_r, column=6, value=r[col_var_abs]).number_format = "+#,##0.00;-#,##0.00;0.00"
        ws_dash.cell(row=curr_r, column=7, value=r["Var. % Recente"] / 100).number_format = "+0.00%;-0.00%;0.00%"
        
        for col_idx in range(2, 8):
            cell = ws_dash.cell(row=curr_r, column=col_idx)
            cell.font = font_cell
            cell.border = border_light
            if col_idx in [6, 7]:
                cell.fill = fill_red
                cell.font = font_red
        curr_r += 1

    # Tabela: Maiores Reduções Recentes
    curr_r += 2
    ws_dash.cell(row=curr_r, column=2, value=f"✅ TOP REDUÇÕES NO ÚLTIMO MÊS ({penultimo_mes} ➔ {ultimo_mes})").font = font_section
    curr_r += 1

    for idx, h in enumerate(headers_resumo, start=2):
        cell = ws_dash.cell(row=curr_r, column=idx, value=h)
        cell.fill = fill_header
        cell.font = font_head
        cell.alignment = Alignment(horizontal="center", vertical="center")

    curr_r += 1
    for _, r in maiores_quedas.iterrows():
        ws_dash.cell(row=curr_r, column=2, value=r["CÓDIGO"]).alignment = Alignment(horizontal="center")
        ws_dash.cell(row=curr_r, column=3, value=r["EVENTO"])
        ws_dash.cell(row=curr_r, column=4, value=r[penultimo_mes]).number_format = "#,##0.00"
        ws_dash.cell(row=curr_r, column=5, value=r[ultimo_mes]).number_format = "#,##0.00"
        ws_dash.cell(row=curr_r, column=6, value=r[col_var_abs]).number_format = "+#,##0.00;-#,##0.00;0.00"
        ws_dash.cell(row=curr_r, column=7, value=r["Var. % Recente"] / 100).number_format = "+0.00%;-0.00%;0.00%"
        
        for col_idx in range(2, 8):
            cell = ws_dash.cell(row=curr_r, column=col_idx)
            cell.font = font_cell
            cell.border = border_light
            if col_idx in [6, 7]:
                cell.fill = fill_green
                cell.font = font_green
        curr_r += 1

    ws_dash.column_dimensions['A'].width = 3
    ws_dash.column_dimensions['B'].width = 12
    ws_dash.column_dimensions['C'].width = 32
    ws_dash.column_dimensions['D'].width = 16
    ws_dash.column_dimensions['E'].width = 16
    ws_dash.column_dimensions['F'].width = 22
    ws_dash.column_dimensions['G'].width = 16

    # ABA 2: MATRIZ COMPLETA
    ws_comp = wb.create_sheet(title="Matriz Comparativa Completa")
    headers_comp = list(df_pivot.columns)

    for c_idx, h in enumerate(headers_comp, start=1):
        cell = ws_comp.cell(row=1, column=c_idx, value=h)
        cell.fill = fill_header
        cell.font = font_head
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, r_data in enumerate(df_pivot.values, start=2):
        for c_idx, val in enumerate(r_data, start=1):
            cell = ws_comp.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.border = border_light
            
            col_name = headers_comp[c_idx - 1]
            if col_name in meses_nomes or col_name == col_var_abs:
                cell.number_format = "#,##0.00"
            elif col_name == "Var. % Recente":
                cell.value = val / 100.0
                cell.number_format = "+0.00%;-0.00%;0.00%"
            elif col_name == "CÓDIGO":
                cell.alignment = Alignment(horizontal="center")
            elif col_name == "Tendência Recente":
                cell.alignment = Alignment(horizontal="center")
                # SETAS FORMATADAS E COLORIDAS NO EXCEL
                if "Aumentou" in str(val):
                    cell.value = "▲ Aumentou"
                    cell.fill = fill_red
                    cell.font = font_red
                elif "Diminuiu" in str(val):
                    cell.value = "▼ Diminuiu"
                    cell.fill = fill_green
                    cell.font = font_green
                else:
                    cell.value = "▬ Estável"
                    cell.fill = fill_gray
                    cell.font = font_gray

    for col in ws_comp.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_comp.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # ABA 3: BASE BRUTA CONSOLIDADA
    ws_base = wb.create_sheet(title="Base Consolidada Globus")
    headers_base = list(df_detalhado_geral.columns)

    for c_idx, h in enumerate(headers_base, start=1):
        cell = ws_base.cell(row=1, column=c_idx, value=h)
        cell.fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
        cell.font = font_head

    for r_idx, r_data in enumerate(df_detalhado_geral.values, start=2):
        for c_idx, val in enumerate(r_data, start=1):
            cell = ws_base.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.border = border_light
            if headers_base[c_idx-1] == "REFERENCIA":
                cell.number_format = "#,##0.00"

    wb.save(buffer)

    st.sidebar.markdown("---")
    st.sidebar.success("🎉 Relatório Executivo MoM Gerado!")
    st.sidebar.download_button(
        label="📥 Baixar Relatório Executivo MoM (.XLSX)",
        data=buffer.getvalue(),
        file_name=f"Relatorio_MoM_Globus_{penultimo_mes}_vs_{ultimo_mes}.xlsx",
        mime="application/vnd.ms-excel"
    )

elif len(arquivos_carregados) == 1:
    st.warning("⚠️ Selecione pelo menos 2 meses para ativar o relatório comparativo.")
else:
    st.info("👈 Utilize o menu lateral para selecionar a quantidade de meses (2 a 5) e faça o upload dos arquivos TXT.")
